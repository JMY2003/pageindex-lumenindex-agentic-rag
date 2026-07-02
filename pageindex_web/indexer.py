import os
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

import fitz

from .schemas import DocumentStructureRow, extract_json_array, validate_model_list
from .storage import CACHE_SCHEMA_VERSION, INDEX_VERSION, PDF_LOCAL_INDEX_VERSION

Progress = Callable[[int, str], None]
LARGE_NODE_MIN_PAGES = int(os.getenv("PAGEINDEX_LARGE_NODE_MIN_PAGES", "12"))
LARGE_NODE_MIN_CHARS = int(os.getenv("PAGEINDEX_LARGE_NODE_MIN_CHARS", "24000"))
MARKDOWN_MIN_NODE_TOKENS = int(os.getenv("PAGEINDEX_MARKDOWN_MIN_NODE_TOKENS", "80"))


def _clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _page_summary(text: str, limit: int = 360) -> str:
    text = _clean_text(text)
    return text[:limit] + ("..." if len(text) > limit else "")


def _estimate_tokens(text: str) -> int:
    text = text or ""
    cjk = len(re.findall(r"[\u3400-\u9fff]", text))
    latin_words = len(re.findall(r"[A-Za-z0-9_]+", text))
    other = max(0, len(text) - cjk)
    return cjk + latin_words + other // 6


def _node(node_id: str, title: str, start: int, end: int, summary: str = "", children: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    return {
        "node_id": node_id,
        "title": title or "Untitled",
        "start_index": int(start),
        "end_index": int(end),
        "pages": [int(start), int(end)],
        "summary": summary,
        "nodes": children or [],
    }


def _assign_ids(nodes: List[Dict[str, Any]], start: int = 1) -> int:
    current = start
    for node in nodes:
        node["node_id"] = str(current).zfill(4)
        current += 1
        current = _assign_ids(node.get("nodes", []), current)
    return current


def _entries_to_tree(entries: List[Dict[str, Any]], pages: List[Dict[str, Any]], max_end: Optional[int] = None) -> List[Dict[str, Any]]:
    if not entries:
        return []
    boundary_end = min(len(pages), int(max_end or len(pages)))
    stack: List[Tuple[int, Dict[str, Any]]] = []
    roots: List[Dict[str, Any]] = []
    for i, entry in enumerate(entries):
        next_same_or_higher = next((e["start"] for e in entries[i + 1 :] if e["level"] <= entry["level"]), boundary_end + 1)
        end = max(entry["start"], min(boundary_end, next_same_or_higher - 1))
        text = " ".join(p["content"] for p in pages[entry["start"] - 1 : end])
        item = _node("", entry["title"], entry["start"], end, _page_summary(text))
        item["approximate"] = bool(entry.get("approximate", False))
        if "top_ratio" in entry:
            item["top_ratio"] = entry["top_ratio"]
        while stack and stack[-1][0] >= entry["level"]:
            stack.pop()
        if stack:
            stack[-1][1].setdefault("nodes", []).append(item)
        else:
            roots.append(item)
        stack.append((entry["level"], item))
    _assign_ids(roots)
    return roots


def _section_text(pages: List[Dict[str, Any]], start: int, end: int) -> str:
    start = max(1, int(start))
    end = min(len(pages), max(start, int(end)))
    return " ".join(p.get("content", "") for p in pages[start - 1 : end])


def _bookmark_entries(doc: fitz.Document, pages: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    raw = doc.get_toc(simple=False)
    entries: List[Dict[str, Any]] = []
    prev_level = 0
    for item in raw:
        if len(item) < 3:
            continue
        level, title, page_num = int(item[0]), _clean_text(item[1]), int(item[2])
        if not title or page_num < 1:
            continue
        if prev_level and level > prev_level + 1:
            level = prev_level + 1
        prev_level = level
        top_ratio = 0.0
        dest = item[3] if len(item) > 3 else {}
        try:
            point = dest.get("to") if isinstance(dest, dict) else None
            if point is not None and 1 <= page_num <= len(doc):
                top_ratio = max(0.0, min(1.0, float(point.y) / max(1.0, doc[page_num - 1].rect.height)))
        except Exception:
            top_ratio = 0.0
        entries.append({
            "level": max(1, min(6, level)),
            "title": title,
            "start": max(1, min(page_num, len(pages))),
            "top_ratio": round(top_ratio, 4),
            "approximate": False,
        })
    return entries


def _line_texts(page: fitz.Page) -> List[Dict[str, Any]]:
    lines = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            spans = [s for s in line.get("spans", []) if _clean_text(s.get("text", ""))]
            if not spans:
                continue
            text = _clean_text(" ".join(s.get("text", "") for s in spans))
            if not text:
                continue
            bbox = [
                min(float(s.get("bbox", [0, 0, 0, 0])[0]) for s in spans),
                min(float(s.get("bbox", [0, 0, 0, 0])[1]) for s in spans),
                max(float(s.get("bbox", [0, 0, 0, 0])[2]) for s in spans),
                max(float(s.get("bbox", [0, 0, 0, 0])[3]) for s in spans),
            ]
            lines.append({
                "text": text,
                "bbox": bbox,
                "size": max(float(s.get("size", 0)) for s in spans),
                "bold": any("bold" in str(s.get("font", "")).lower() for s in spans),
            })
    return lines


def _repeated_headers_footers(doc: fitz.Document) -> set[str]:
    counts: Dict[str, int] = {}
    total = len(doc)
    for page in doc:
        height = page.rect.height
        seen = set()
        for line in _line_texts(page):
            y0, y1 = line["bbox"][1], line["bbox"][3]
            if y0 < height * 0.12 or y1 > height * 0.88:
                text = _clean_text(line["text"]).lower()
                if 3 <= len(text) <= 90:
                    seen.add(text)
        for text in seen:
            counts[text] = counts.get(text, 0) + 1
    threshold = max(3, int(total * 0.35))
    return {text for text, count in counts.items() if count >= threshold}


def _is_toc_page(lines: List[Dict[str, Any]]) -> bool:
    if not lines:
        return False
    text = "\n".join(line["text"] for line in lines[:80])
    lower = text.lower()
    dot_leaders = sum(1 for line in lines if re.search(r"\.{4,}\s*\d{1,4}$", line["text"]))
    numbered = sum(1 for line in lines if re.search(r"\s\d{1,4}$", line["text"]))
    has_label = any(token in lower for token in ("table of contents", "contents", "\u76ee\u5f55", "\u76ee\u6b21"))
    return has_label and (dot_leaders >= 3 or numbered >= 8)


def _extract_pdf_pages(doc: fitz.Document) -> Tuple[List[Dict[str, Any]], List[int]]:
    repeated = _repeated_headers_footers(doc)
    pages = []
    toc_pages = []
    for idx, page in enumerate(doc, 1):
        lines = _line_texts(page)
        if _is_toc_page(lines):
            toc_pages.append(idx)
        content_lines = []
        for line in lines:
            text_key = _clean_text(line["text"]).lower()
            if text_key in repeated:
                continue
            content_lines.append(line["text"])
        pages.append({"page": idx, "content": "\n".join(content_lines)})
    return pages, toc_pages


def _pdf_heading_entries(doc: fitz.Document, toc_pages: List[int]) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    font_sizes: List[float] = []
    for page_number, page in enumerate(doc, 1):
        if page_number in toc_pages:
            continue
        raw_lines = _line_texts(page)
        merged: List[Dict[str, Any]] = []
        for line in raw_lines:
            text = line["text"]
            if merged and abs(line["size"] - merged[-1]["size"]) < 0.7 and 0 <= line["bbox"][1] - merged[-1]["bbox"][3] < 10 and len(merged[-1]["text"]) + len(text) < 140:
                merged[-1]["text"] = _clean_text(merged[-1]["text"] + " " + text)
                merged[-1]["bbox"][3] = line["bbox"][3]
                merged[-1]["bold"] = merged[-1]["bold"] or line["bold"]
                continue
            merged.append(dict(line))
        for line in merged:
            text = line["text"]
            if len(text) < 3 or len(text) > 140:
                continue
            size = line["size"]
            font_sizes.append(size)
            bold = line["bold"]
            y0 = line["bbox"][1]
            numbered = bool(re.match(r"^(\d+(\.\d+)*|\u7b2c[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341\u767e\u5343\u4e070-9]+[\u7ae0\u8282\u7bc7\u90e8]|[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341]+[\u3001.])\s*", text))
            decorative = text.isupper() and len(text.split()) <= 3 and len(text) <= 24
            candidates.append({"page": page_number, "title": text, "size": size, "bold": bold, "y0": y0, "numbered": numbered, "decorative": decorative})

    if not candidates:
        return []
    baseline = sorted(font_sizes)[max(0, int(len(font_sizes) * 0.75) - 1)] if font_sizes else 11
    headings = [
        c
        for c in candidates
        if not c["decorative"] and (c["size"] >= baseline + 1.2 or c["bold"] or c["numbered"]) and c["y0"] > 35
    ]
    deduped: List[Dict[str, Any]] = []
    seen = set()
    for c in headings:
        key = (c["page"], c["title"].lower())
        if key in seen:
            continue
        seen.add(key)
        deduped.append(c)
    deduped = deduped[:160]
    if len(deduped) < 2:
        return []

    entries = []
    for heading in deduped:
        entries.append({
            "level": _heading_level(heading["title"], float(heading["size"]), baseline),
            "title": heading["title"],
            "start": heading["page"],
            "approximate": True,
        })
    return entries


def _detect_pdf_headings(doc: fitz.Document, pages: List[Dict[str, Any]], toc_pages: List[int]) -> List[Dict[str, Any]]:
    return _entries_to_tree(_pdf_heading_entries(doc, toc_pages), pages)


def _heading_level(title: str, size: float, baseline: float) -> int:
    text = title.strip()
    if re.match(r"^(part|chapter|appendix)\b", text, re.I):
        return 1
    if re.match(r"^\u7b2c[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341\u767e\u5343\u4e070-9]+[\u7ae0\u8282\u7bc7\u90e8]", text):
        return 1
    match = re.match(r"^(\d+(?:\.\d+){0,5})\b", text)
    if match:
        return max(1, min(6, match.group(1).count(".") + 1))
    if size >= baseline + 4:
        return 1
    if size >= baseline + 2.4:
        return 2
    return 3


def _page_chunks(pages: List[Dict[str, Any]], chunk_size: int = 3) -> List[Dict[str, Any]]:
    roots: List[Dict[str, Any]] = []
    for start in range(1, len(pages) + 1, chunk_size):
        end = min(len(pages), start + chunk_size - 1)
        text = " ".join(p["content"] for p in pages[start - 1 : end])
        roots.append(_node("", f"Pages {start}-{end}", start, end, _page_summary(text)))
    _assign_ids(roots)
    return roots


def _should_refine_node(node: Dict[str, Any], pages: List[Dict[str, Any]]) -> bool:
    if node.get("nodes"):
        return False
    start, end = (node.get("pages") or [node.get("start_index", 1), node.get("end_index", 1)])[:2]
    page_span = int(end) - int(start) + 1
    if page_span >= LARGE_NODE_MIN_PAGES:
        return True
    return len(_section_text(pages, int(start), int(end))) >= LARGE_NODE_MIN_CHARS


def _entries_for_range(entries: List[Dict[str, Any]], start: int, end: int, parent_title: str) -> List[Dict[str, Any]]:
    parent_key = _clean_text(parent_title).lower()
    rows: List[Dict[str, Any]] = []
    for entry in entries:
        title = _clean_text(entry.get("title", ""))
        page = int(entry.get("start") or 0)
        if not title or page < start or page > end:
            continue
        if page == start and title.lower() == parent_key:
            continue
        rows.append(dict(entry))
    if len(rows) < 2:
        return []
    min_level = min(int(row.get("level") or 1) for row in rows)
    normalized = []
    for row in rows:
        row["level"] = max(1, int(row.get("level") or 1) - min_level + 1)
        normalized.append(row)
    return normalized


def _refine_large_pdf_nodes(nodes: List[Dict[str, Any]], pages: List[Dict[str, Any]], heading_entries: List[Dict[str, Any]]) -> int:
    refined = 0
    for node in nodes:
        refined += _refine_large_pdf_nodes(node.get("nodes", []), pages, heading_entries)
        if not _should_refine_node(node, pages):
            continue
        start, end = (node.get("pages") or [node.get("start_index", 1), node.get("end_index", 1)])[:2]
        range_entries = _entries_for_range(heading_entries, int(start), int(end), str(node.get("title", "")))
        children = _entries_to_tree(range_entries, pages, max_end=int(end))
        children = [
            child for child in children
            if int(child.get("start_index", 0)) >= int(start) and int(child.get("end_index", 0)) <= int(end)
        ]
        if len(children) >= 2:
            node["nodes"] = children
            node["refined_from_layout"] = True
            refined += 1
    if refined:
        _assign_ids(nodes)
    return refined


def _extract_json_array(text: str) -> List[Dict[str, Any]]:
    rows = validate_model_list(DocumentStructureRow, extract_json_array(text), limit=120)
    return [row.model_dump() for row in rows]


def _llm_structure_from_pages(kind: str, pages: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    try:
        from .llm import completion, has_api_key
        from .prompts import get_prompt
        from .settings import SettingsStore

        settings = SettingsStore(Path(__file__).resolve().parent.parent / "app_settings.json").load(include_secret=True)
        if not has_api_key(settings):
            return []
        sample = "\n\n".join(f"Page {page['page']}:\n{(page.get('content') or '')[:1800]}" for page in pages[:12])
        prompt = [
            {"role": "system", "content": get_prompt("doc_structure.system")},
            {"role": "user", "content": get_prompt("doc_structure.user", kind=kind, page_count=len(pages), sample=sample)},
        ]
        response = completion(settings, prompt, max_tokens=1200, temperature=0)
        rows = _extract_json_array(response.text)
    except Exception:
        return []
    roots: List[Dict[str, Any]] = []
    stack: List[Tuple[int, Dict[str, Any]]] = []
    for row in rows[:120]:
        title = _clean_text(str(row.get("title", "")))
        if not title:
            continue
        level = max(1, min(6, int(row.get("level") or 1)))
        start = max(1, min(len(pages), int(row.get("start_page") or 1)))
        end = max(start, min(len(pages), int(row.get("end_page") or start)))
        text = " ".join(p["content"] for p in pages[start - 1 : end])
        item = _node("", title, start, end, _page_summary(text), [])
        while stack and stack[-1][0] >= level:
            stack.pop()
        if stack:
            stack[-1][1].setdefault("nodes", []).append(item)
        else:
            roots.append(item)
        stack.append((level, item))
    if roots:
        _assign_ids(roots)
    return roots


def _split_text_pages(text: str, chars_per_page: int = 3500) -> List[Dict[str, Any]]:
    chunks = []
    current = []
    size = 0
    for paragraph in re.split(r"(\n\s*\n)", text):
        if size + len(paragraph) > chars_per_page and current:
            chunks.append("".join(current).strip())
            current = []
            size = 0
        current.append(paragraph)
        size += len(paragraph)
    if current:
        chunks.append("".join(current).strip())
    return [{"page": i + 1, "content": chunk} for i, chunk in enumerate(chunks or [text])]


def _markdown_heading_rows(text: str) -> List[Dict[str, Any]]:
    lines = text.splitlines()
    heading_rows = []
    char_pos = 0
    in_fence = False
    fence_marker = ""
    for line_no, line in enumerate(lines, 1):
        stripped = line.strip()
        fence = re.match(r"^(```+|~~~+)", stripped)
        if fence:
            marker = fence.group(1)[:3]
            if not in_fence:
                in_fence = True
                fence_marker = marker
            elif marker == fence_marker:
                in_fence = False
                fence_marker = ""
            char_pos += len(line) + 1
            continue
        match = None if in_fence or line.startswith(("    ", "\t")) else re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            heading_rows.append({
                "level": len(match.group(1)),
                "title": _clean_text(match.group(2)),
                "line": line_no,
                "char": char_pos,
            })
        char_pos += len(line) + 1
    return heading_rows


def _prune_private_fields(nodes: List[Dict[str, Any]]) -> None:
    for node in nodes:
        node.pop("_token_count", None)
        node.pop("_char_span", None)
        _prune_private_fields(node.get("nodes", []))


def _thin_markdown_nodes(nodes: List[Dict[str, Any]], min_tokens: int = MARKDOWN_MIN_NODE_TOKENS) -> int:
    removed = 0
    for node in nodes:
        removed += _thin_markdown_nodes(node.get("nodes", []), min_tokens)
        kept_children = []
        for child in node.get("nodes", []):
            child_tokens = int(child.get("_token_count") or 0)
            if child_tokens and child_tokens < min_tokens and not child.get("nodes"):
                removed += 1
                continue
            kept_children.append(child)
        node["nodes"] = kept_children
    return removed


def _index_markdown_text(text: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], str]:
    heading_rows = _markdown_heading_rows(text)

    pages = _split_text_pages(text)
    if len(heading_rows) < 2:
        return _page_chunks(pages, chunk_size=1), pages, "markdown_chunks"

    def page_for_char(pos: int) -> int:
        total = 0
        for page in pages:
            total += len(page["content"]) + 1
            if pos <= total:
                return page["page"]
        return pages[-1]["page"]

    roots: List[Dict[str, Any]] = []
    stack: List[Tuple[int, Dict[str, Any]]] = []
    for i, row in enumerate(heading_rows):
        next_row = next((r for r in heading_rows[i + 1 :] if r["level"] <= row["level"]), None)
        start_char = row["char"]
        end_char = next_row["char"] if next_row else len(text)
        section_text = text[start_char:end_char]
        start_page = page_for_char(start_char)
        end_page = max(start_page, page_for_char(max(start_char, end_char - 1)))
        item = _node("", row["title"], start_page, end_page, _page_summary(section_text), [])
        item["line_num"] = row["line"]
        item["_token_count"] = _estimate_tokens(section_text)
        item["_char_span"] = [start_char, end_char]
        while stack and stack[-1][0] >= row["level"]:
            stack.pop()
        if stack:
            stack[-1][1].setdefault("nodes", []).append(item)
        else:
            roots.append(item)
        stack.append((row["level"], item))
    thinned = _thin_markdown_nodes(roots) if len(heading_rows) >= 24 else 0
    _prune_private_fields(roots)
    _assign_ids(roots)
    strategy = "markdown_headings_thinned" if thinned else "markdown_headings"
    return roots, pages, strategy


def index_pdf(path: Path, progress: Progress) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    progress(12, "Reading PDF pages")
    doc = fitz.open(path)
    pages, toc_pages = _extract_pdf_pages(doc)
    progress(45, f"Read {len(doc)} pages; filtering repeated headers and footers")

    progress(52, "Analyzing document outline")
    bookmark_entries = _bookmark_entries(doc, pages)
    structure = _entries_to_tree(bookmark_entries, pages)
    strategy = "pdf_bookmarks"
    layout_entries = _pdf_heading_entries(doc, toc_pages)
    if len(structure) < 2:
        structure = _entries_to_tree(layout_entries, pages)
        strategy = "pdf_layout_headings"
    if len(structure) < 2:
        structure = _page_chunks(pages)
        strategy = "pdf_page_chunks"
    refined_nodes = 0
    if layout_entries and structure and strategy in {"pdf_bookmarks", "pdf_layout_headings"}:
        refined_nodes = _refine_large_pdf_nodes(structure, pages, layout_entries)
        if refined_nodes:
            strategy = f"{strategy}_refined"
    doc.close()
    progress(78, f"Building index tree: {strategy}")
    metadata = {
        "structure": structure,
        "summary": _page_summary(" ".join(p["content"] for p in pages[: min(3, len(pages))]), 600),
        "version": INDEX_VERSION,
        "cache_schema_version": CACHE_SCHEMA_VERSION,
        "pdf_local_index_version": PDF_LOCAL_INDEX_VERSION,
        "index_strategy": strategy,
        "refined_large_nodes": refined_nodes,
        "toc_pages": toc_pages,
        "page_count": len(pages),
    }
    return metadata, pages


def _docx_pages_with_word_com(path: Path) -> Optional[List[Dict[str, Any]]]:
    if os.name != "nt":
        return None
    script = r'''
param([string]$Path)
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$doc = $word.Documents.Open($Path, $false, $true)
$items = @()
$pageCount = $doc.ComputeStatistics(2)
$imageIndex = 1
for ($i = 1; $i -le $pageCount; $i++) {
  $startRange = $doc.GoTo(1, 1, $i)
  $endPos = $doc.Content.End
  if ($i -lt $pageCount) {
    $nextRange = $doc.GoTo(1, 1, ($i + 1))
    $endPos = $nextRange.Start
  }
  $range = $doc.Range($startRange.Start, $endPos)
  $text = $range.Text -replace "`r", "`n" -replace "`a", ""
  $placeholders = @()
  foreach ($shape in $doc.InlineShapes) {
    try {
      if ($shape.Range.Information(3) -eq $i) {
        $placeholders += "[image $imageIndex]"
        $imageIndex += 1
      }
    } catch {}
  }
  foreach ($shape in $doc.Shapes) {
    try {
      if ($shape.Anchor.Information(3) -eq $i) {
        $placeholders += "[image $imageIndex]"
        $imageIndex += 1
      }
    } catch {}
  }
  if ($placeholders.Count -gt 0) {
    $text = $text + "`n" + ($placeholders -join "`n")
  }
  $items += [PSCustomObject]@{ page=$i; text=$text }
}
$doc.Close($false)
$word.Quit()
$items | ConvertTo-Json -Depth 4
'''
    try:
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script, str(path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=180,
        )
        data = json.loads(completed.stdout or "[]")
    except Exception:
        return None
    if isinstance(data, dict):
        data = [data]
    pages: Dict[int, List[str]] = {}
    for item in data:
        page = int(item.get("page") or 1)
        text = _clean_text(item.get("text", ""))
        if text:
            pages.setdefault(page, []).append(text)
    return [{"page": page, "content": "\n".join(lines)} for page, lines in sorted(pages.items())] or None


def _docx_pages_from_xml(path: Path) -> Optional[List[Dict[str, Any]]]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        root = ET.fromstring(xml)
    except Exception:
        return None
    pages: List[List[str]] = [[]]
    for para in root.findall(".//w:body/w:p", ns):
        parts: List[str] = []
        for elem in para.iter():
            if elem.tag == f"{{{ns['w']}}}lastRenderedPageBreak":
                if parts:
                    pages[-1].append("".join(parts))
                    parts = []
                pages.append([])
            elif elem.tag == f"{{{ns['w']}}}br" and elem.attrib.get(f"{{{ns['w']}}}type") == "page":
                if parts:
                    pages[-1].append("".join(parts))
                    parts = []
                pages.append([])
            elif elem.tag == f"{{{ns['w']}}}t" and elem.text:
                parts.append(elem.text)
        if parts:
            pages[-1].append("".join(parts))
    output = [{"page": idx + 1, "content": "\n".join(lines)} for idx, lines in enumerate(pages) if "\n".join(lines).strip()]
    return output or None


def _docx_paragraphs(path: Path) -> List[Dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    rows = []
    current_page = 1
    for para in doc.paragraphs:
        text = _clean_text(para.text)
        if not text:
            continue
        style = para.style.name if para.style else ""
        heading_match = re.match(r"Heading\s+([1-6])", style, re.I)
        page_breaks = sum(1 for run in para.runs for br in run._element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br") if br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page")
        rows.append({"text": text, "style": style, "level": int(heading_match.group(1)) if heading_match else 0, "page": current_page})
        current_page += page_breaks
        if len(" ".join(r["text"] for r in rows if r["page"] == current_page)) > 2600:
            current_page += 1
    return rows


def index_docx(path: Path, progress: Progress) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    progress(20, "Parsing Word paragraphs")
    com_pages = _docx_pages_with_word_com(path)
    rows = _docx_paragraphs(path)
    page_count = max([r["page"] for r in rows], default=1)
    if com_pages:
        pages = com_pages
        page_count = max(page["page"] for page in pages)
        paging_strategy = "word_com"
    else:
        xml_pages = _docx_pages_from_xml(path)
        if xml_pages:
            pages = xml_pages
            page_count = max(page["page"] for page in pages)
            paging_strategy = "xml_rendered_page_breaks"
        else:
            pages = []
            for page in range(1, page_count + 1):
                pages.append({"page": page, "content": "\n".join(r["text"] for r in rows if r["page"] == page)})
            paging_strategy = "xml_fallback"

    progress(55, "Reading heading styles")
    heading_rows = [r for r in rows if r["level"]]
    roots: List[Dict[str, Any]] = []
    stack: List[Tuple[int, Dict[str, Any]]] = []
    page_content = {page["page"]: page.get("content", "") for page in pages}
    for i, row in enumerate(heading_rows):
        next_same = next((r["page"] for r in heading_rows[i + 1 :] if r["level"] <= row["level"]), page_count + 1)
        item = _node("", row["text"], row["page"], max(row["page"], next_same - 1), _page_summary(page_content.get(row["page"], "")))
        while stack and stack[-1][0] >= row["level"]:
            stack.pop()
        if stack:
            stack[-1][1].setdefault("nodes", []).append(item)
        else:
            roots.append(item)
        stack.append((row["level"], item))

    if len(roots) < 3:
        roots = _llm_structure_from_pages("DOCX", pages)
        if len(roots) >= 2:
            strategy = "docx_llm_toc"
        else:
            roots = _page_chunks(pages, chunk_size=2)
            strategy = "docx_page_chunks"
    else:
        strategy = "docx_headings"
        _assign_ids(roots)

    metadata = {
        "structure": roots,
        "summary": _page_summary(" ".join(p["content"] for p in pages[: min(3, len(pages))]), 600),
        "version": INDEX_VERSION,
        "cache_schema_version": CACHE_SCHEMA_VERSION,
        "index_strategy": strategy,
        "paging_strategy": paging_strategy,
        "page_count": len(pages),
    }
    return metadata, pages


def _convert_doc_to_docx(path: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice/soffice was not found, so .doc files cannot be converted. Install LibreOffice and retry.")
    tmpdir = Path(tempfile.mkdtemp(prefix="pageindex_doc_"))
    subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", str(tmpdir), str(path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    converted = tmpdir / (path.stem + ".docx")
    if not converted.exists():
        raise RuntimeError("DOC to DOCX conversion failed: no converted output was generated.")
    return converted


def index_document(path: Path, progress: Progress) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return index_pdf(path, progress)
    if suffix == ".docx":
        return index_docx(path, progress)
    if suffix == ".doc":
        progress(10, "Converting DOC to DOCX")
        return index_docx(_convert_doc_to_docx(path), progress)
    if suffix in {".md", ".markdown"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        structure, pages, strategy = _index_markdown_text(text)
        metadata = {
            "structure": structure,
            "summary": _page_summary(text, 600),
            "version": INDEX_VERSION,
            "cache_schema_version": CACHE_SCHEMA_VERSION,
            "index_strategy": strategy,
            "page_count": len(pages),
        }
        return metadata, pages
    raise ValueError(f"Unsupported file type: {suffix}")
